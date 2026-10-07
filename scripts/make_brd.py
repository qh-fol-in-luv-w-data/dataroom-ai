"""Fill the DAIT BRD Word template with the AI Data Room project, keeping its layout.

Copies the template, replaces the text of every paragraph / table cell in place (the first
run's formatting is kept), and grows or shrinks tables by cloning a body row.
Usage: .venv/bin/python scripts/make_brd.py TEMPLATE.docx OUT.docx
"""
from __future__ import annotations

import copy
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

TITLE = ("DỰ ÁN: AI CORE CHO DATA ROOM IPO", " TRÊN SHAREPOINT")
DATE = ("Ngày lập: 07/10/2026", " - ", "Phiên bản: v1.0")
HEADER = "DAIT – DATA ROOM IPO | BUSINESS REQUIREMENT DOCUMENT | AI CORE CHO DATA ROOM"

INFO = [
    ("Đơn vị nghiệp vụ", "Chủ Data Room IPO (Deal team / Ban IPO); Quản trị tenant Microsoft 365 và Purview"),
    ("Đơn vị triển khai", "Ban DAIT - Phát triển ứng dụng – Team AI"),
    ("Mục tiêu giai đoạn", "Cổng đi tiếp sau 9 phép thử tích hợp ngày 23/10/2026; Go-live Bản 1 dự kiến 26/02/2027 (tuần 20)"),
    ("Phạm vi chính", "AI core on-prem tại Việt Nam trả lời câu hỏi của nhà đầu tư và nhân viên trên tài liệu Data Room "
                      "SharePoint, đúng quyền SharePoint, mọi ý có trích dẫn nguyên văn đã kiểm chứng"),
    ("Tình trạng", "Dự thảo – Kế hoạch tổng v0.6 và bản demo local đã chạy, có số đo; đang chờ Chủ Data Room, "
                   "Quản trị tenant và BOD xác nhận các Business Rule"),
]

H = {
    "scope": "MỤC TIÊU & PHẠM VI (SCOPE)",
    "pipeline": "PIPELINE NGHIỆP VỤ CỐT LÕI",
    "br": "DANH SÁCH YÊU CẦU NGHIỆP VỤ (BUSINESS REQUIREMENTS)",
    "rules": "BUSINESS RULE CẦN CHỦ DATA ROOM / BOD XÁC NHẬN",
    "inputs": "CÁC BÊN CẦN CUNG CẤP (INPUTS)",
    "uat": "TIÊU CHÍ UAT/PILOT ĐỀ XUẤT",
    "ms": "MỐC TRIỂN KHAI (MILESTONES)",
    "matrix": "MA TRẬN ĐỐI CHIẾU 28 YÊU CẦU GỐC (KẾ HOẠCH BẢO MẬT AI DATA ROOM)",
    "refs": "TÀI LIỆU THAM KHẢO CHÍNH",
}

OBJECTIVE = ("Mục tiêu: Xây AI core chạy on-prem tại Việt Nam cho Data Room IPO trên SharePoint Online. Nhà đầu tư và nhân "
             "viên hỏi bằng tiếng Việt (có dấu hoặc không dấu) hoặc tiếng Anh; AI chỉ trả lời từ tài liệu người hỏi được mở "
             "trên SharePoint và được phép dùng AI; mọi ý có trích dẫn nguyên văn đã kiểm chứng; không rò rỉ chéo giữa các "
             "nhóm nhà đầu tư; trên các câu AI trả lời, độ đúng ≥ 99% với độ tin cậy 95%.")
IN_SCOPE = ("Phạm vi In-scope: Bắt đầu từ khi tài liệu đã nằm trên site Data Room SharePoint (đã gán AI_Use và nhãn) → "
            "Đồng bộ quyền và nạp tài liệu (an toàn file, đọc/OCR, cắt đoạn, lập chỉ mục) → Hỏi đáp qua giao diện chat "
            "(đăng nhập Entra) với 4 cổng quyền → Trả lời có trích dẫn, link mở file trên SharePoint → Nhật ký và báo cáo.")
OUT_SCOPE = ("Phạm vi Out-of-scope: Không xây Data Room, trình xem, watermark hay chức năng cấp quyền (thuộc SharePoint, "
             "Entra, Purview do Chủ Data Room và Quản trị tenant thực hiện). Không lưu file gốc on-prem. Không dùng Copilot "
             "hay agent của Microsoft cho nhà đầu tư. Câu hỏi đếm/tổng và kho bản ghi thuộc Bản 2; không tư vấn đầu tư.")

PIPE_INTRO = "Quy trình hệ thống tuân theo luồng xử lý:"
PIPE_B0 = ("B0 – Chuẩn bị Data Room: Chủ Data Room đưa tài liệu lên SharePoint, gán AI_Use và nhãn; Quản trị tenant mời "
           "nhà đầu tư thành khách Entra B2B và cấp access package theo nhóm. AI core bắt đầu xử lý khi tài liệu xuất hiện "
           "trong delta của Microsoft Graph.")
PIPE = ["B1 Đồng bộ quyền & delta → B2 Nạp tài liệu (an toàn file, đọc/OCR) → B3 Lập chỉ mục theo phân vùng",
        "B4 Hỏi & kiểm 4 cổng quyền → B5 Tìm kiếm, rerank & kiểm tra trực tiếp SharePoint",
        "B6 Sinh câu trả lời & kiểm chứng trích dẫn → B7 Trả lời một khối & ghi nhật ký"]

BR = [
    ("BR-01", "Đăng nhập & vai trò AI",
     "Nhà đầu tư (khách Entra B2B) và nhân viên đăng nhập bằng Entra (MFA, Conditional Access). Quyền dùng AI cấp qua app "
     "role Investor.AI / Investor.AI.Summarize / Staff.AI / Reviewer theo access package của từng nhóm.",
     "Không tạo tài khoản riêng trong AI core. Người không có app role AI không hỏi được, vẫn mở file trên SharePoint bình thường."),
    ("BR-02", "Đồng bộ quyền & Thu hồi",
     "Đọc quyền SharePoint/Entra qua Graph delta (60 giây), dựng bản chép quyền. Gỡ khỏi nhóm, chuyển nhóm, offboard có "
     "hiệu lực ở AI ngay cả khi SharePoint còn trễ; token còn hạn của người bị offboard cũng bị chặn.",
     "Chỉ quyền đọc (Sites.Selected). Nguồn quyền ngừng cập nhật quá ngưỡng → từ chối mọi câu hỏi (đóng an toàn)."),
    ("BR-03", "Nạp tài liệu đa định dạng & OCR",
     "Nạp Word, Excel, PowerPoint, PDF có chữ, PDF scan/ảnh, email, HTML, CSV, TXT. Bản scan được làm sạch (nắn ảnh chụp, "
     "xoay, chỉnh nghiêng, phóng to) rồi OCR bằng PaddleOCR-VL; con số được bộ đọc thứ hai đối chiếu. Đổi phiên bản thì "
     "đoạn của bản cũ bị vô hiệu ngay.",
     "File AI_Use = Excluded, nhãn mã hóa, có mã độc → không nạp, ghi lý do. Không giữ file gốc on-prem. Trang có số "
     "chưa xác nhận hoặc chữ viết tay → hàng chờ người duyệt."),
    ("BR-04", "Hỏi đáp đúng quyền (4 cổng)",
     "Chỉ trả lời từ tài liệu người hỏi mở được trên SharePoint và được phép AI: RLS trên bản chép quyền + kiểm tra trực "
     "tiếp với SharePoint bằng token người hỏi + tập thu hồi + danh sách chặn.",
     "Bắt buộc 0 rò rỉ chéo giữa các nhóm nhà đầu tư. Từ chối trung tính, không tiết lộ tài liệu ngoài quyền có tồn tại."),
    ("BR-05", "Trích dẫn & Kiểm chứng",
     "Mọi ý có trích dẫn nguyên văn; con số trong ý phải có trong câu trích; NLI làm tín hiệu phụ; số từ bản scan chưa xác "
     "nhận gắn nhãn \"chưa chắc\"; link mở đúng file và vị trí trên SharePoint.",
     "Câu trả lời giữ ở server tới khi kiểm xong, gửi một khối. Không trả downloadUrl hay file gốc; ảnh cắt trang scan chỉ "
     "khi file không chặn tải."),
    ("BR-06", "Bảo mật đầu vào/đầu ra",
     "Chuẩn hóa câu hỏi, chặn prompt injection (luật tiếng Việt/Anh + bộ phân loại), che dữ liệu cá nhân (CCCD, di động, "
     "email), hạn mức theo người dùng; tài liệu luôn được coi là dữ liệu không đáng tin.",
     "Không tư vấn đầu tư, không dự báo. Bộ lọc chặn nhầm > 1% trên câu hỏi thật thì chuyển sang chỉ ghi log."),
    ("BR-07", "Phạm vi truy xuất & Minh bạch",
     "Người dùng xem được AI truy xuất được những file nào trong phạm vi của mình, file nào bị loại và lý do (Excluded, "
     "mã hóa, đang nạp, AI đang chặn do thu hồi).",
     "Không bao giờ hiện file ngoài phạm vi SharePoint của người hỏi."),
    ("BR-08", "Audit Log & Kiểm toán",
     "Ghi mọi câu hỏi, câu trả lời, file và phiên bản đã trích, từ chối, chặn, thay đổi quyền; chuỗi hash chống sửa; neo "
     "WORM trên SharePoint; đối chiếu với nhật ký Purview.",
     "Phục vụ kiểm toán IPO. Thời hạn lưu do luật sư chốt; dữ liệu cá nhân được che trong nhật ký."),
    ("BR-09", "Vận hành & Dashboard",
     "Theo dõi độ tươi nguồn đồng bộ, model, chỉ mục, hàng chờ duyệt; runbook RB-1…RB-7; công tắc tắt AI theo site hoặc "
     "toàn hệ thống.",
     "Hypercare 4 tuần, ca trực có tên ≥ 2 người."),
    ("BR-10", "Câu hỏi đếm/tổng (Bản 2)",
     "Trả lời câu hỏi đếm và tổng trên kho bản ghi trích từ tài liệu, kiểm tra trực tiếp toàn bộ file nguồn.",
     "Không bao giờ trả một tổng thiếu; ngoài phạm vi Bản 1."),
]

RULES = [
    ("1", "AI có được tổng hợp câu trả lời cho nhà đầu tư (bật Investor.AI.Summarize) hay chỉ tìm và trích nguyên văn?",
     "BR-01, BR-05, phạm vi PoC và cổng độ chính xác.", "Chờ BOD xác nhận (trước tuần 6)"),
    ("2", "Tách mỗi nhóm nhà đầu tư một site và bật RAC không?",
     "BR-02, BR-04, chỉ mục theo phân vùng; bắt buộc nếu phép thử 3 thất bại.", "Chờ Chủ Data Room (tuần 2)"),
    ("3", "Gán AI_Use cho toàn bộ file đang có: mặc định Excluded, thư mục nào được Allowed?",
     "BR-03, độ phủ AI, lịch nạp production.", "Chờ Chủ Data Room (tuần 6)"),
    ("4", "Có cho trích câu từ tài liệu chỉ-xem (chặn tải) và cho hiển thị ảnh cắt trang scan không?",
     "BR-05, chính sách hiển thị.", "Chờ Chủ Data Room (tuần 2)"),
    ("5", "Thay \"chính xác 100% / bảo mật tuyệt đối\" bằng: độ đúng ≥ 99% (tin cậy 95%), từ chối nhầm ≤ 15%, 0 rò rỉ.",
     "Cổng go-live, tiêu chí UAT, truyền thông với nhà đầu tư.", "Chờ BOD xác nhận"),
    ("6", "Bộ lọc prompt injection: chặn hay chỉ ghi log; hạn mức câu hỏi của mỗi người dùng.",
     "BR-06, trải nghiệm nhà đầu tư.", "Đề xuất: chặn khi tin cậy cao; 8 câu/phút, 60 câu/giờ"),
    ("7", "Thời hạn lưu nhật ký AI và site neo WORM.", "BR-08, Preservation Lock.", "Chờ luật sư (tuần 12)"),
    ("8", "Hạn mức ngân sách GPU on-prem.", "Toàn bộ dự án, lịch go-live.", "Chờ BOD (tuần 0)"),
]

INPUTS = [
    ("1", "Site thử Data Room, ≥ 10 tài khoản khách thử (B2B), cấp quyền Sites.Selected, đồng ý quyền directory-reader.",
     "Quản trị tenant", "Cuối tuần 1 (16/10/2026) – để chạy 9 phép thử tích hợp"),
    ("2", "Danh sách người dùng OTP kiểu cũ trên site Data Room và chuyển sang khách B2B.",
     "Quản trị tenant", "Trước 31/10/2026 (Microsoft ngừng OTP kiểu cũ)"),
    ("3", "Bộ câu hỏi chuẩn: 150 câu phát triển + ≥ 500 câu giữ riêng có đáp án (hai người chấm).",
     "Trưởng deal team cử người nghiệp vụ", "Tập phát triển 30/10/2026; tập giữ riêng 18/12/2026"),
    ("4", "300 trang scan mẫu (dấu đỏ, bảng, chữ ký, viết tay) + 1.000 ô số chép lại làm đáp án OCR.",
     "Người nghiệp vụ", "Trước 23/10/2026 (PoC OCR)"),
    ("5", "Gán AI_Use, nhãn \"IPO Data Room\", DLP chặn Copilot, Restricted Content Discovery trên production.",
     "Chủ Data Room, Quản trị Purview", "Tuần 6 (20/11/2026)"),
    ("6", "Ý kiến pháp lý bằng văn bản; thời hạn lưu nhật ký.", "Luật sư / Pháp chế",
     "Thời hạn lưu 31/12/2026; ý kiến pháp lý trước 08/02/2027"),
    ("7", "Hạn mức ngân sách GPU; cấu hình GPU chọn sau PoC.", "BOD", "Ngân sách 09/10/2026; cấu hình 30/10/2026"),
]

UAT = [
    ("UAT-01", "Nhà đầu tư hỏi bộ câu giữ riêng (≥ 500 câu, đủ loại nguồn, có dấu/không dấu).",
     "Độ đúng ≥ 99% (tin cậy 95%) trên câu đã trả lời; từ chối nhầm ≤ 15%; 100% trích dẫn nguyên văn hợp lệ."),
    ("UAT-02", "Nhà đầu tư Nhóm B hỏi nội dung chỉ Nhóm A hoặc site Nội bộ có.",
     "Từ chối trung tính, không lộ nội dung hay tên file; 0 lỗi trên ma trận ≥ 3.000 phép thử + kiểm thử xâm nhập."),
    ("UAT-03", "Gỡ nhà đầu tư khỏi nhóm, offboard, đổi AI_Use sang Excluded, thay file bằng bản đã biên tập.",
     "Nhóm và offboard: AI chặn ≤ 2 phút; câu kế tiếp không dùng bản cũ hay file bị loại."),
    ("UAT-04", "Ngắt nguồn đồng bộ quyền (Entra/SharePoint).", "Hệ thống từ chối mọi câu hỏi trong ≤ 3 phút (đóng an toàn)."),
    ("UAT-05", "Hỏi về bản scan khó (xoay, ảnh chụp, fax, dấu đỏ đè số).",
     "Con số khớp đáp án người chép; số chưa xác nhận luôn có nhãn \"chưa chắc\" và ảnh vùng gốc (nếu được phép)."),
    ("UAT-06", "Bắn bộ tấn công prompt injection đa ngôn ngữ (≥ 300 câu) và câu hỏi thật.",
     "0 mẫu làm lộ dữ liệu; chặn nhầm ≤ 1% trên câu hỏi thật."),
    ("UAT-07", "Đo tải giờ cao điểm, có đồng bộ chạy song song.",
     "0 file bị loại vì throttling; p95 thời gian trả lời trong ngưỡng chốt ở PoC."),
]

MS_INTRO = ("Dự án dự kiến 20 tuần cho Bản 1 + 4 tuần hypercare (đội 4–5 người), tính từ tuần 1 = 12/10/2026; "
            "ngày go-live gắn với lịch IPO.")
MILESTONES = [
    ("07/10/2026", "Kế hoạch tổng v0.6 & Demo local",
     "Demo 4 cổng quyền, OCR bản scan khó, tìm kiếm 1 triệu đoạn, lớp bảo mật đầu vào đã chạy và có số đo. (Đã thực hiện)"),
    ("09/10/2026", "Tuần 0 – Ngân sách", "BOD chốt hạn mức ngân sách GPU; người bảo trợ cấp điều hành được chỉ định."),
    ("23/10/2026", "9 phép thử tích hợp M365", "Cổng đi tiếp/dừng. Phép thử 3 thất bại → chuyển chế độ dự phòng."),
    ("30/10/2026", "PoC độ chính xác & Cổng BOD",
     "So sánh LLM, OCR bake-off, chọn cấu hình GPU. Hạn chuyển OTP sang B2B: 31/10/2026."),
    ("04/12/2026", "Nạp tài liệu & Quyền", "Connector, bản chép quyền, tập thu hồi, nạp, chỉ mục phân vùng."),
    ("08/01/2027", "Luồng truy vấn", "Gateway, kiểm tra trực tiếp, giao diện qua DMZ, kiểm chứng L5, nhật ký, neo WORM."),
    ("22/01/2027", "Kiểm thử & Kiểm thử xâm nhập",
     "Ma trận rò rỉ, thu hồi, đo tải, red-team, pentest độc lập, chấm bộ giữ riêng, tập runbook."),
    ("05/02/2027", "Sửa & Kiểm lại", "Sửa mọi phát hiện, pentest lại. Lưu ý Tết Nguyên đán (~06/02/2027): có thể lùi 1 tuần."),
    ("19/02/2027", "Thử nội bộ", "Nhân viên và cố vấn dùng với quyền production; test rò rỉ theo nhóm hằng ngày."),
    ("26/02/2027", "Go-live Bản 1", "Mở đồng thời cho các nhóm nhà đầu tư. Một cổng trượt → dời ít nhất 2 tuần."),
    ("26/03/2027", "Kết thúc Hypercare", "Ca trực ≥ 2 người; sau đó bắt đầu Bản 2 (8–10 tuần)."),
]

MATRIX = [
    ("Rò rỉ dữ liệu chéo", "AI không trả lời thông tin mật của nhà đầu tư này cho nhà đầu tư khác", "BR-04", "Must"),
    ("Prompt injection trực tiếp", "Chặn câu hỏi ép bỏ quy tắc, lộ cấu hình, vượt quyền", "BR-06", "Must"),
    ("Injection gián tiếp qua tài liệu", "Lệnh cài trong tài liệu chỉ được coi là dữ liệu, có gắn cờ", "BR-03, BR-06", "Must"),
    ("Đa ngôn ngữ", "Câu hỏi và tấn công tiếng Việt, tiếng Anh, trộn lẫn, ký tự ẩn", "BR-06", "Must"),
    ("Phân quyền trước truy xuất", "AI không tự quyết quyền; lọc quyền trước và trong lúc truy xuất", "BR-04", "Must"),
    ("Quyền AI_QUERY", "Cho phép AI đọc tài liệu để trả lời", "BR-01, BR-04", "Must"),
    ("Quyền AI_SUMMARIZE", "Cho phép AI tóm tắt, tổng hợp", "BR-01 / Pending Decision", "Pending Decision"),
    ("Quyền AI_CITATION", "Trích dẫn nguồn (phiên bản, trang)", "BR-05", "Must"),
    ("Lớp 1 – SSO, MFA, IAM", "Định danh và phân quyền người dùng", "BR-01", "Must"),
    ("Lớp 2 – Input security", "Chuẩn hóa Unicode, nhận diện ngôn ngữ, lọc PII, chặn ghi đè chỉ dẫn", "BR-06", "Must"),
    ("Lớp 3 – Document security", "Quét mã độc, phát hiện nội dung ẩn trong PDF/Excel", "BR-03", "Must"),
    ("Lớp 4 – Permission-aware RAG", "Lọc theo quyền, quản lý phiên bản tài liệu", "BR-02, BR-04", "Must"),
    ("Lớp 5 – Output security", "Chống hallucination, DLP, lưu vết", "BR-05, BR-06, BR-08", "Must"),
    ("Metadata quyền trên vector DB", "Thay bằng bản chép quyền + RLS + kiểm tra trực tiếp SharePoint", "BR-02, BR-04", "Must"),
    ("Tài liệu/quyền hết hạn", "Quyền có hạn và access review hằng tháng", "BR-02", "Should"),
    ("Guardrail tự host", "Bộ lọc tự xây trên bge-m3 (Prompt Guard 2 bị khóa, không có tiếng Việt)", "BR-06", "Must"),
    ("Red-team đa ngôn ngữ", "Bộ tấn công giữ riêng ≥ 300 câu + câu bẫy hợp lệ", "BR-06, UAT-06", "Must"),
    ("Prompt có thẻ <untrusted_document>", "Đóng gói tài liệu với ranh giới ngẫu nhiên mỗi lần hỏi", "BR-06", "Must"),
    ("Grounding", "Chỉ trả lời từ tài liệu; không có thì từ chối trung tính", "BR-05", "Must"),
    ("OCR PDF scan", "PaddleOCR-VL + làm sạch trang (thay Tesseract/Textract)", "BR-03", "Must"),
    ("Excel giữ ngữ cảnh", "Đọc trực tiếp xlsx, giữ đơn vị, kỳ, sheet ẩn", "BR-03", "Must"),
    ("Làm sạch ký tự ẩn", "Bỏ zero-width, ký tự giả dạng ở khâu nạp và ở câu hỏi", "BR-03, BR-06", "Must"),
    ("Cắt đoạn theo cấu trúc", "Không cắt giữa bảng tài chính", "BR-03", "Should"),
    ("Audit log riêng cho AI", "Ai hỏi gì, lúc nào, file và phiên bản nào", "BR-08", "Must"),
    ("API security, quản lý secret", "Chỉ kết nối chiều ra, WAF tại DMZ, kho khóa riêng", "BR-09", "Must"),
    ("Data residency", "LLM on-prem tại Việt Nam (thay proxy tới OpenAI)", "BR-04, BR-09", "Must"),
    ("Chống cạn tài nguyên", "Hạn mức theo người, trần số câu xử lý đồng thời", "BR-06, BR-09", "Must"),
    ("Số liệu tài chính chính xác", "Con số phải có trong câu trích; định nghĩa đo ≥ 99% (tin cậy 95%)", "BR-05", "Must"),
]

REFS = ["• Kế hoạch tổng: AI core cho Data Room trên SharePoint (IPO), bản v0.6 ngày 07/10/2026.",
        "• Kế hoạch triển khai và bảo mật hệ thống AI cho Data Room (IPO) – tài liệu tham khảo.",
        "• Bản demo local và số đo ngày 07/10/2026 (README, logs của dự án dataroom-ai-demo)."]


# ------------------------------------------------------------------ helpers
def set_par(p: Paragraph, text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_runs(p: Paragraph, parts) -> None:
    runs = p.runs
    for i, t in enumerate(parts):
        if i < len(runs):
            runs[i].text = t
    for r in runs[len(parts):]:
        r._r.getparent().remove(r._r)


def set_cell(cell, text: str) -> None:
    ps = cell.paragraphs
    set_par(ps[0], text)
    for p in ps[1:]:
        p._p.getparent().remove(p._p)


def fill_table(t: Table, rows, header: tuple | None = None) -> None:
    body = t.rows[1:]
    tpl = copy.deepcopy(body[0]._tr)
    while len(t.rows) - 1 < len(rows):
        t.rows[-1]._tr.addnext(copy.deepcopy(tpl))
    while len(t.rows) - 1 > len(rows):
        t._tbl.remove(t.rows[-1]._tr)
    if header:
        for c, h in zip(t.rows[0].cells, header):
            set_cell(c, h)
    for r, vals in zip(t.rows[1:], rows):
        for c, v in zip(r.cells, vals):
            set_cell(c, v)


def main(src: str, out: str) -> int:
    d = docx.Document(src)
    els = list(d.element.body.iterchildren())
    P = lambda i: Paragraph(els[i], d)
    T = lambda i: Table(els[i], d)

    set_runs(P(5), TITLE)
    set_runs(P(9), DATE)
    info = T(12)
    for r, (k, v) in zip(info.rows, INFO):
        set_cell(r.cells[0], k)
        set_cell(r.cells[1], v)

    set_par(P(15), H["scope"])
    set_par(P(16), OBJECTIVE)
    set_par(P(17), IN_SCOPE)
    set_par(P(18), OUT_SCOPE)
    set_par(P(19), H["pipeline"])
    set_par(P(20), PIPE_INTRO)
    set_par(P(21), PIPE_B0)
    for i, line in zip((22, 23, 24), PIPE):
        set_par(P(i), line)
    set_par(P(25), H["br"])
    fill_table(T(26), BR)
    set_par(P(27), H["rules"])
    fill_table(T(28), RULES)
    set_par(P(29), H["inputs"])
    fill_table(T(30), INPUTS)
    set_par(P(31), H["uat"])
    fill_table(T(32), UAT)
    set_par(P(33), H["ms"])
    set_par(P(34), MS_INTRO)
    fill_table(T(35), MILESTONES)
    set_par(P(36), H["matrix"])
    fill_table(T(37), [(str(i), *row) for i, row in enumerate(MATRIX, 1)],
               header=("STT", "Yêu cầu gốc", "Nội dung yêu cầu gốc", "BR/Use Case liên quan", "Phân loại"))
    set_par(P(38), H["refs"])
    set_par(P(39), REFS[0])
    set_par(P(40), REFS[1])
    last = els[40]
    for text in REFS[2:]:
        new = copy.deepcopy(last)
        last.addnext(new)
        set_par(Paragraph(new, d), text)
        last = new

    for s in d.sections:
        for p in s.header.paragraphs:
            if p.text.strip():
                set_par(p, HEADER)
    d.core_properties.title = "BRD – AI Core cho Data Room IPO trên SharePoint"
    d.save(out)
    print("Đã lưu", out)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
